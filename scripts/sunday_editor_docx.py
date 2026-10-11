"""Editor-first DOCX projection of a validated Sunday worksheet.

This is a presentation layer only. The unchanged original TXT is the
provenance/approval artifact; this document is a provisional editing copy,
not a publication-ready IPR Brief or a substitute for original-source checks.
"""
from __future__ import annotations

import io
import re
from urllib.parse import urlsplit

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.shared import Inches, Pt

from scripts.sunday_editor_readable import (
    APPENDIX, BEGIN, PROSE_TITLES, editable_sections,
)

RECORD_PREFIX = "SOURCE RECORD IDS:"
EXTERNAL_PREFIX = "EXTERNAL SOURCE IDS:"
RECORD_LINE = re.compile(r"^Record ([0-9]+) \|")
EXTERNAL_LINE = re.compile(r"^External source ([A-Za-z0-9_-]+) \|")
BOLD = re.compile(r"\*\*([^\n*]+)\*\*")


def _references(packet):
    """Return ONLY cited publisher URLs, never raw bodies/synopses/hashes."""
    manuscript = packet.split(BEGIN, 1)[1].split(APPENDIX, 1)[0]
    appendix = packet.split(APPENDIX, 1)[1]
    cited = []
    for line in manuscript.splitlines():
        kind = None
        if line.startswith(RECORD_PREFIX):
            kind = "Record"
            values = line[len(RECORD_PREFIX):]
        elif line.startswith(EXTERNAL_PREFIX):
            kind = "External"
            values = line[len(EXTERNAL_PREFIX):]
        if kind:
            for raw in values.split(","):
                identifier = raw.strip()
                if not re.fullmatch(r"[A-Za-z0-9_-]+", identifier):
                    raise ValueError("invalid cited source identifier in worksheet")
                key = (kind, identifier)
                if key not in cited:
                    cited.append(key)

    catalog = {}
    current = None
    for line in appendix.splitlines():
        record = RECORD_LINE.match(line)
        external = EXTERNAL_LINE.match(line)
        if record:
            current = ("Record", record.group(1))
        elif external:
            current = ("External", external.group(1))
        elif line.startswith("URL: ") and current is not None:
            url = line[5:].strip()
            parts = urlsplit(url)
            if (parts.scheme == "https" and parts.hostname
                    and not parts.username and not parts.password):
                if current in catalog and catalog[current] != url:
                    raise ValueError("duplicate source identity with conflicting URL")
                catalog[current] = url
            current = None
    missing = [name for name in cited if name not in catalog]
    if missing:
        raise ValueError("cited source URL absent from immutable worksheet")
    return [(kind, identifier, catalog[(kind, identifier)])
            for kind, identifier in cited]


def _rich_paragraph(paragraph, value):
    """Preserve authored text and optional **bold** without interpreting HTML."""
    clean = " ".join(value.splitlines()).strip()
    offset = 0
    for match in BOLD.finditer(clean):
        if match.start() > offset:
            paragraph.add_run(clean[offset:match.start()])
        paragraph.add_run(match.group(1)).bold = True
        offset = match.end()
    if offset < len(clean):
        paragraph.add_run(clean[offset:])


def _hyperlink(paragraph, label, target):
    relation = paragraph.part.relate_to(
        target, RELATIONSHIP_TYPE.HYPERLINK, is_external=True,
    )
    element = OxmlElement("w:hyperlink")
    element.set(qn("r:id"), relation)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    style = OxmlElement("w:rStyle")
    style.set(qn("w:val"), "Hyperlink")
    props.append(style)
    run.append(props)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    element.append(run)
    paragraph._p.append(element)


def render_editor_docx(worksheet):
    """Return .docx bytes only after the source-text boundary validator passes."""
    sections = editable_sections(worksheet)
    sources = _references(worksheet)
    doc = Document()
    doc.core_properties.title = "Indo-Pacific Record — Working Brief"
    doc.core_properties.author = "Indo-Pacific Record"

    section = doc.sections[0]
    section.top_margin = Inches(.82)
    section.bottom_margin = Inches(.78)
    section.left_margin = Inches(.85)
    section.right_margin = Inches(.85)

    normal = doc.styles["Normal"]
    normal.font.name = "Georgia"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.24
    for name in ("Title", "Heading 1", "Heading 2"):
        style = doc.styles[name]
        style.font.name = "Aptos"
        style.font.color.rgb = None

    label = doc.add_paragraph()
    label.alignment = WD_ALIGN_PARAGRAPH.CENTER
    label.add_run("INDO-PACIFIC RECORD / BRIEFS").bold = True

    doc.add_paragraph(
        "Provisional editorial manuscript · For Dylan's language and structure edits only"
    )
    doc.add_paragraph(
        "Ben will independently check the underlying sources and approve any publication."
    )
    for key, value in sections:
        if key == "WORKING TITLE":
            p = doc.add_paragraph(style="Title")
            _rich_paragraph(p, value)
        elif key == "DEK":
            p = doc.add_paragraph(style="Subtitle")
            _rich_paragraph(p, value)
        else:
            doc.add_heading(PROSE_TITLES[key], level=2)
            for block in re.split(r"\n\s*\n", value):
                if block.strip():
                    p = doc.add_paragraph()
                    _rich_paragraph(p, block)

    if sources:
        doc.add_page_break()
        doc.add_heading("Source links for context", level=2)
        doc.add_paragraph(
            "These are links cited by the draft, not independently verified "
            "publication approvals. Please flag any factual questions; Ben "
            "will check original wording, dates and final citations."
        )
        for kind, identifier, url in sources:
            p = doc.add_paragraph()
            _hyperlink(p, kind + " " + identifier, url)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("UNNUMBERED DRAFT · NOT APPROVED FOR PUBLICATION")

    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()
