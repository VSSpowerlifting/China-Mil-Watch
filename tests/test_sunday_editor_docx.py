"""Offline contracts: editable DOCX never includes technical source bodies."""
from __future__ import annotations

from io import BytesIO
import unittest
from zipfile import ZipFile

from docx import Document

from scripts.sunday_editor_docx import render_editor_docx
from scripts.sunday_editor_readable import ReadableEditorError, editable_sections
from tests.test_sunday_editor_readable import reviewed_packet


class WordHandoffTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.original = reviewed_packet()

    def test_word_has_complete_editable_manuscript_and_no_appendix(self):
        output = render_editor_docx(self.original)
        self.assertTrue(output.startswith(b"PK"))
        with ZipFile(BytesIO(output)) as archive:
            self.assertIn("word/document.xml", archive.namelist())
        document = Document(BytesIO(output))
        combined = "\n".join(p.text for p in document.paragraphs)
        sections = editable_sections(self.original)
        self.assertIn(sections[0][1], combined)
        for key, raw in sections:
            if key not in ("WORKING TITLE", "DEK"):
                self.assertIn(raw.splitlines()[0].split("**")[0].strip(), combined)
        self.assertIn("UNNUMBERED DRAFT", document.sections[0].footer.paragraphs[0].text)
        for forbidden in (
            "SOURCE RECORD IDS:", "EXTERNAL SOURCE IDS:",
            "Pinned shadow commit:", "Version/content SHA-256:",
            "MODEL-AVAILABLE OFFICIAL SOURCE RESEARCH",
            "=== SOURCE APPENDIX", "COVERAGE SNAPSHOT",
            "MANUSCRIPT SOURCE USE", "Research synopsis (not approved):",
        ):
            self.assertNotIn(forbidden, combined)

    def test_word_only_links_to_cited_official_source_urls(self):
        doc = Document(BytesIO(render_editor_docx(self.original)))
        targets = [
            relation.target_ref for relation in doc.part.rels.values()
            if relation.reltype.endswith("/hyperlink")
        ]
        self.assertGreater(len(targets), 0)
        self.assertEqual(len(targets), len(set(targets)))
        self.assertTrue(all(value.startswith("https://") for value in targets))

    def test_invalid_manuscript_does_not_render(self):
        mutated = self.original.replace("=== EDITABLE MANUSCRIPT ===", "WRONG", 1)
        with self.assertRaises(ReadableEditorError):
            render_editor_docx(mutated)
        broken = self.original.replace("URL: https://", "URL: javascript:", 1)
        # If a cited URL is replaced by an unsafe scheme, DOCX creation fails closed.
        with self.assertRaises(ValueError):
            render_editor_docx(broken)


if __name__ == "__main__":
    unittest.main()
